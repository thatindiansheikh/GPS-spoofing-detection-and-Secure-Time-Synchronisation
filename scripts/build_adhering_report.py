import os
import sys
import re
from pathlib import Path

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

BASE_DIR = Path(r"Y:\Final yr project\Honours\gps-spoof-timesync")
REPORT_DIR = BASE_DIR / "report"
FIG_DIR = REPORT_DIR / "generated_figures"
MEDIA_DIR = REPORT_DIR / "extracted_media"

AU_LOGO = str(MEDIA_DIR / "image2.png")
SVCE_LOGO = str(MEDIA_DIR / "image3.png")

FIG_1_1 = str(FIG_DIR / "figure_1_1_proposed_methodology.png")
FIG_3_1 = str(FIG_DIR / "figure_3_1_system_architecture.png")
FIG_4_1 = str(MEDIA_DIR / "image6.png")
FIG_4_2 = str(MEDIA_DIR / "image7.png")
FIG_5_1 = str(MEDIA_DIR / "image8.png")
FIG_5_2 = str(MEDIA_DIR / "image9.png")
FIG_5_3 = str(MEDIA_DIR / "image10.png")
FIG_6_1 = str(MEDIA_DIR / "image11.png")
FIG_6_2 = str(FIG_DIR / "figure_6_2_flowchart_system_operation.png")
FIG_6_3 = str(FIG_DIR / "figure_6_3_flowchart_decision_logic.png")
FIG_6_4 = str(FIG_DIR / "figure_6_4_holdover_model.png")
FIG_7_1 = str(FIG_DIR / "figure_7_1_deviation_constellation.png")
FIG_7_2 = str(FIG_DIR / "figure_7_2_scenario_timeline.png")
FIG_7_3 = str(FIG_DIR / "figure_7_3_phase_development.png")

def clean(text):
    if not text:
        return ""
    t = str(text)
    for dash in ["\u2013", "\u2014", "\u2212", "\u2010", "\u2011", "\u2012", "\u2015", "-"]:
        t = t.replace(dash, " ")
    t = re.sub(r" +", " ", t)
    return t.strip()

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False, color_rgb=(0,0,0)):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{name}" w:hAnsi="{name}" w:cs="{name}"/>')
    rPr.append(rFonts)

def set_section_margins(section):
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

def add_page_number_to_section(section, num_format="decimal", start_num=None):
    footer = section.footer
    footer.is_linked_to_previous = False
    p_footer = footer.paragraphs[0]
    p_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_footer.text = ""
    
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag.endswith("pgNumType"):
            sectPr.remove(child)
            
    attrs = f'w:fmt="{num_format}"'
    if start_num is not None:
        attrs += f' w:start="{start_num}"'
    pgNumType = parse_xml(f'<w:pgNumType {nsdecls("w")} {attrs}/>')
    sectPr.append(pgNumType)
    
    run = p_footer.add_run()
    set_font(run, name="Times New Roman", size=11, bold=False)
    
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve">PAGE</w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)

def p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False, italic=False, before=0, after=6, spacing=1.5, first_line=0.5, keep_with_next=False):
    c_text = clean(text)
    par = doc.add_paragraph()
    par.alignment = align
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = spacing
    par.paragraph_format.first_line_indent = Inches(first_line)
    par.paragraph_format.keep_with_next = keep_with_next
    if c_text:
        run = par.add_run(c_text)
        set_font(run, size=size, bold=bold, italic=italic)
    return par

def h1(doc, text, before=18, after=12):
    return p(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, before=before, after=after, first_line=0.0, keep_with_next=True)

def h2(doc, text, before=14, after=6):
    return p(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=True, before=before, after=after, first_line=0.0, keep_with_next=True)

def h3(doc, text, before=10, after=4):
    return p(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, size=12, bold=True, italic=True, before=before, after=after, first_line=0.0, keep_with_next=True)

def bullet(doc, title, desc, after=4):
    c_title = clean(title)
    c_desc = clean(desc)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.left_indent = Inches(0.25)
    par.paragraph_format.first_line_indent = Inches(0)
    par.paragraph_format.space_before = Pt(2)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.3
    
    r_sym = par.add_run(" ")
    set_font(r_sym, size=11, bold=True)
    if c_title:
        r_title = par.add_run(c_title + ": ")
        set_font(r_title, size=11, bold=True)
    r_desc = par.add_run(c_desc)
    set_font(r_desc, size=11, bold=False)
    return par

def add_toc_line(doc, col1, col2, col3, bold=False, size=11, before=1, after=2):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.first_line_indent = Inches(0)
    
    pPr = par._p.get_or_add_pPr()
    tabs = parse_xml(f'''
        <w:tabs {nsdecls("w")}>
            <w:tab w:val="left" w:pos="1440"/>
            <w:tab w:val="right" w:leader="dot" w:pos="8200"/>
        </w:tabs>
    ''')
    pPr.append(tabs)
    
    r1 = par.add_run(clean(col1))
    set_font(r1, size=size, bold=bold)
    
    par.add_run("\t")
    
    r2 = par.add_run(clean(col2))
    set_font(r2, size=size, bold=bold)
    
    par.add_run("\t")
    
    r3 = par.add_run(clean(col3))
    set_font(r3, size=size, bold=bold)
    return par

def add_two_col_line(doc, col1, col2, bold=False, size=11, before=2, after=2, tab_pos=2.2):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.first_line_indent = Inches(0)
    
    pos_dxa = int(tab_pos * 1440)
    pPr = par._p.get_or_add_pPr()
    tabs = parse_xml(f'''
        <w:tabs {nsdecls("w")}>
            <w:tab w:val="left" w:pos="{pos_dxa}"/>
        </w:tabs>
    ''')
    pPr.append(tabs)
    
    r1 = par.add_run(clean(col1))
    set_font(r1, size=size, bold=True)
    
    par.add_run("\t")
    
    r2 = par.add_run(clean(col2))
    set_font(r2, size=size, bold=bold)
    return par

def insert_image(doc, img_path, width_in=5.2, caption_no="", caption_title="", space_before=8, space_after=4):
    c_no = clean(caption_no)
    c_title = clean(caption_title)
    
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    par.paragraph_format.keep_with_next = True
    
    if os.path.exists(img_path):
        run = par.add_run()
        run.add_picture(img_path, width=Inches(width_in))
    else:
        r = par.add_run(f"[ Image File Not Found: {clean(os.path.basename(img_path))} ]")
        set_font(r, size=10, italic=True)
        
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(space_after)
    
    rc1 = cap.add_run(f"Figure {c_no} ")
    set_font(rc1, size=11, bold=True)
    rc2 = cap.add_run(c_title)
    set_font(rc2, size=11, bold=False)
    return par

def table_header(row, col_titles, widths=None):
    for idx, title in enumerate(col_titles):
        cell = row.cells[idx]
        if widths and idx < len(widths):
            cell.width = Inches(widths[idx])
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E2E8F0"/>')
        tcPr.append(shd)
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="333333"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="333333"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="333333"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="333333"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(4)
        r = cp.add_run(clean(title))
        set_font(r, size=10, bold=True)

def table_row(row, col_values, widths=None, align_center=False):
    for idx, val in enumerate(col_values):
        cell = row.cells[idx]
        if widths and idx < len(widths):
            cell.width = Inches(widths[idx])
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center or idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
        cp.paragraph_format.space_before = Pt(3)
        cp.paragraph_format.space_after = Pt(3)
        r = cp.add_run(clean(str(val)))
        set_font(r, size=9.5, bold=False)

print("Helper functions defined successfully!")
