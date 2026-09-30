import docx

doc = docx.Document(r"report\Project_Phase-I_Report.docx")

print("=== SECTIONS IN Project_Phase-I_Report.docx ===")
for i, s in enumerate(doc.sections):
    print(f"Section {i+1}: start_type={s.start_type}, diff_first={s.different_first_page_header_footer}")
    header = s.header
    footer = s.footer
    print(f"  Header text: {[p.text for p in header.paragraphs if p.text]}")
    print(f"  Footer text: {[p.text for p in footer.paragraphs if p.text]}")
    # Inspect footer XML for page number field
    for p in footer.paragraphs:
        print(f"  Footer XML: {p._p.xml}")

print("\n=== SEARCHING FOR TABLE OF CONTENTS in Project_Phase-I_Report.docx ===")
for i, p in enumerate(doc.paragraphs):
    if "TABLE OF CONTENTS" in p.text.upper() or "CHAPTER NO." in p.text.upper():
        print(f"P {i}: {p.text}")
        for j in range(i, min(i+40, len(doc.paragraphs))):
            print(f"  {j}: {doc.paragraphs[j].text}")
        break
